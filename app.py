from flask import Flask, render_template, request, redirect, url_for, session, flash, send_file
from werkzeug.security import generate_password_hash, check_password_hash
from functools import wraps
import sqlite3
from datetime import datetime, timedelta
import csv
import io
import os

app = Flask(__name__)
app.secret_key = 'tool-storage-secret-key-2024'


# ==================== БАЗА ДАННЫХ ====================

def init_db():
    conn = sqlite3.connect('tools.db')
    c = conn.cursor()

    c.execute('''CREATE TABLE IF NOT EXISTS users (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        username TEXT UNIQUE NOT NULL,
        password_hash TEXT NOT NULL,
        role TEXT NOT NULL DEFAULT 'storekeeper'
    )''')

    c.execute('''CREATE TABLE IF NOT EXISTS employees (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        fullname TEXT NOT NULL,
        department TEXT,
        is_active INTEGER DEFAULT 1
    )''')

    c.execute('''CREATE TABLE IF NOT EXISTS tools (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        name TEXT NOT NULL,
        inventory_number TEXT UNIQUE,
        description TEXT,
        status TEXT DEFAULT 'ok'
    )''')

    c.execute('''CREATE TABLE IF NOT EXISTS transactions (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        tool_id INTEGER NOT NULL,
        employee_id INTEGER NOT NULL,
        issued_by INTEGER NOT NULL,
        issue_date TEXT NOT NULL,
        return_date TEXT,
        returned_to INTEGER,
        status TEXT DEFAULT 'issued',
        FOREIGN KEY (tool_id) REFERENCES tools(id),
        FOREIGN KEY (employee_id) REFERENCES employees(id),
        FOREIGN KEY (issued_by) REFERENCES users(id),
        FOREIGN KEY (returned_to) REFERENCES users(id)
    )''')

    c.execute('''CREATE TABLE IF NOT EXISTS repairs (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        tool_id INTEGER NOT NULL,
        sent_date TEXT NOT NULL,
        issue_description TEXT,
        return_date TEXT,
        repair_cost REAL DEFAULT 0,
        notes TEXT,
        FOREIGN KEY (tool_id) REFERENCES tools(id)
    )''')

    c.execute('''CREATE TABLE IF NOT EXISTS settings (
        key TEXT PRIMARY KEY,
        value TEXT NOT NULL
    )''')

    c.execute("INSERT OR IGNORE INTO settings (key, value) VALUES ('deadline_days', '14')")

    # Админ по умолчанию
    c.execute("SELECT id FROM users WHERE username = 'admin'")
    if not c.fetchone():
        c.execute("INSERT INTO users (username, password_hash, role) VALUES (?, ?, ?)",
                  ('admin', generate_password_hash('admin'), 'admin'))

    # Кладовщик по умолчанию
    c.execute("SELECT id FROM users WHERE username = 'storekeeper'")
    if not c.fetchone():
        c.execute("INSERT INTO users (username, password_hash, role) VALUES (?, ?, ?)",
                  ('storekeeper', generate_password_hash('123'), 'storekeeper'))

    # Тестовые сотрудники
    c.execute("SELECT COUNT(*) FROM employees")
    if c.fetchone()[0] == 0:
        employees = [
            ('Иванов Иван Иванович', 'Цех №1'),
            ('Петров Петр Петрович', 'Цех №2'),
            ('Сидоров Сидор Сидорович', 'Склад'),
            ('Кузнецов Алексей Викторович', 'Цех №1'),
            ('Морозова Елена Алексеевна', 'Администрация')
        ]
        c.executemany("INSERT INTO employees (fullname, department) VALUES (?, ?)", employees)

    # Тестовые инструменты (каждый уникальный)
    c.execute("SELECT COUNT(*) FROM tools")
    if c.fetchone()[0] == 0:
        tools = [
            ('Дрель электрическая', 'INV-001', 'Мощная дрель Bosch', 'ok'),
            ('Дрель электрическая', 'INV-002', 'Мощная дрель Bosch', 'ok'),
            ('Шуруповерт Makita', 'INV-003', 'Аккумуляторный 18V', 'ok'),
            ('Шуруповерт Makita', 'INV-004', 'Аккумуляторный 18V', 'ok'),
            ('Набор гаечных ключей', 'INV-005', 'Набор 10-24 мм', 'ok'),
            ('Молоток', 'INV-006', 'Слесарный 500г', 'ok'),
            ('Молоток', 'INV-007', 'Слесарный 500г', 'ok'),
            ('Болгарка', 'INV-008', 'УШМ 125мм', 'ok'),
            ('Лазерный уровень', 'INV-009', 'Высокоточный', 'repair'),
            ('Перфоратор', 'INV-010', 'Тяжелый режим работы', 'ok'),
        ]
        c.executemany("INSERT INTO tools (name, inventory_number, description, status) VALUES (?, ?, ?, ?)", tools)

    conn.commit()
    conn.close()


def get_db():
    conn = sqlite3.connect('tools.db')
    conn.row_factory = sqlite3.Row
    return conn


def get_setting(key, default=None):
    conn = get_db()
    row = conn.execute("SELECT value FROM settings WHERE key = ?", (key,)).fetchone()
    conn.close()
    if row:
        return row['value']
    return default


# ==================== ДЕКОРАТОРЫ ====================

def login_required(f):
    @wraps(f)
    def decorated_function(*args, **kwargs):
        if 'user_id' not in session:
            flash('Пожалуйста, войдите в систему', 'warning')
            return redirect(url_for('login'))
        return f(*args, **kwargs)
    return decorated_function


def admin_required(f):
    @wraps(f)
    def decorated_function(*args, **kwargs):
        if 'user_id' not in session:
            return redirect(url_for('login'))
        if session.get('role') != 'admin':
            flash('Доступ запрещен. Только для администратора', 'danger')
            return redirect(url_for('dashboard'))
        return f(*args, **kwargs)
    return decorated_function


# ==================== АВТОРИЗАЦИЯ ====================

@app.route('/')
def index():
    if 'user_id' in session:
        return redirect(url_for('dashboard'))
    return redirect(url_for('login'))


@app.route('/login', methods=['GET', 'POST'])
def login():
    if request.method == 'POST':
        username = request.form['username']
        password = request.form['password']
        conn = get_db()
        user = conn.execute("SELECT * FROM users WHERE username = ?", (username,)).fetchone()
        conn.close()
        if user and check_password_hash(user['password_hash'], password):
            session['user_id'] = user['id']
            session['username'] = user['username']
            session['role'] = user['role']
            flash(f'Добро пожаловать, {user["username"]}!', 'success')
            return redirect(url_for('dashboard'))
        else:
            flash('Неверное имя пользователя или пароль', 'danger')
    return render_template('login.html')


@app.route('/logout')
def logout():
    session.clear()
    flash('Вы вышли из системы', 'info')
    return redirect(url_for('login'))


# ==================== ПАНЕЛЬ УПРАВЛЕНИЯ ====================

@app.route('/dashboard')
@login_required
def dashboard():
    conn = get_db()
    deadline_days = int(get_setting('deadline_days', '14'))

    total_tools = conn.execute("SELECT COUNT(*) FROM tools").fetchone()[0]
    available_tools = conn.execute("SELECT COUNT(*) FROM tools WHERE status = 'ok'").fetchone()[0]
    issued_count = conn.execute("SELECT COUNT(*) FROM tools WHERE status = 'issued'").fetchone()[0]
    in_repair = conn.execute("SELECT COUNT(*) FROM tools WHERE status = 'repair'").fetchone()[0]

    deadline_date = (datetime.now() - timedelta(days=deadline_days)).strftime('%Y-%m-%d')
    overdue = conn.execute(
        "SELECT COUNT(*) FROM transactions WHERE status = 'issued' AND issue_date < ?",
        (deadline_date,)
    ).fetchone()[0]

    # Просроченные с деталями
    overdue_details = conn.execute('''
        SELECT t.id, t.name as tool_name, t.inventory_number,
               e.fullname as employee_name, tr.issue_date,
               CAST(julianday('now') - julianday(tr.issue_date) AS INTEGER) as days_overdue
        FROM transactions tr
        JOIN tools t ON tr.tool_id = t.id
        JOIN employees e ON tr.employee_id = e.id
        WHERE tr.status = 'issued' AND tr.issue_date < ?
        ORDER BY days_overdue DESC
    ''', (deadline_date,)).fetchall()

    # Последние выдачи
    recent = conn.execute('''
        SELECT t.id, t.name as tool_name, t.inventory_number,
               e.fullname as employee_name, tr.issue_date,
               u.username as issued_by_name,
               CAST(julianday('now') - julianday(tr.issue_date) AS INTEGER) as days_issued
        FROM transactions tr
        JOIN tools t ON tr.tool_id = t.id
        JOIN employees e ON tr.employee_id = e.id
        JOIN users u ON tr.issued_by = u.id
        WHERE tr.status = 'issued'
        ORDER BY tr.issue_date DESC
        LIMIT 10
    ''').fetchall()

    repairs = conn.execute('''
        SELECT r.id, t.name as tool_name, t.inventory_number,
               r.sent_date, r.issue_description
        FROM repairs r
        JOIN tools t ON r.tool_id = t.id
        WHERE r.return_date IS NULL
        ORDER BY r.sent_date DESC
    ''').fetchall()

    conn.close()
    return render_template('dashboard.html',
                         total_tools=total_tools,
                         available_tools=available_tools,
                         issued_count=issued_count,
                         in_repair=in_repair,
                         overdue=overdue,
                         deadline_days=deadline_days,
                         overdue_details=overdue_details,
                         recent=recent,
                         repairs=repairs)

# ==================== ВЫДАЧА ИНСТРУМЕНТА ====================

@app.route('/issue', methods=['GET', 'POST'])
@login_required
def issue_tool():
    conn = get_db()
    if request.method == 'POST':
        tool_id = request.form['tool_id']
        employee_id = request.form['employee_id']
        issue_date = request.form['issue_date'] or datetime.now().strftime('%Y-%m-%d')

        # Проверка статуса инструмента
        tool = conn.execute("SELECT * FROM tools WHERE id = ?", (tool_id,)).fetchone()
        if not tool:
            conn.close()
            flash('Инструмент не найден', 'danger')
            return redirect(url_for('issue_tool'))

        if tool['status'] == 'repair':
            conn.close()
            flash('ОШИБКА: Инструмент находится в ремонте! Выдача невозможна.', 'danger')
            return redirect(url_for('issue_tool'))

        if tool['status'] == 'decommissioned':
            conn.close()
            flash('ОШИБКА: Инструмент списан! Выдача невозможна.', 'danger')
            return redirect(url_for('issue_tool'))

        if tool['status'] == 'issued':
            conn.close()
            flash('ОШИБКА: Инструмент уже выдан другому сотруднику!', 'danger')
            return redirect(url_for('issue_tool'))

        # Создаём запись выдачи
        conn.execute('''
            INSERT INTO transactions (tool_id, employee_id, issued_by, issue_date, status)
            VALUES (?, ?, ?, ?, 'issued')
        ''', (tool_id, employee_id, session['user_id'], issue_date))

        # Меняем статус инструмента на "выдано"
        conn.execute("UPDATE tools SET status = 'issued' WHERE id = ?", (tool_id,))

        conn.commit()
        conn.close()
        flash('Инструмент успешно выдан!', 'success')
        return redirect(url_for('dashboard'))

    # GET - показываем форму
    tools = conn.execute("SELECT * FROM tools WHERE status = 'ok' ORDER BY inventory_number").fetchall()
    employees = conn.execute("SELECT * FROM employees WHERE is_active = 1 ORDER BY fullname").fetchall()

    conn.close()
    return render_template('issue.html', tools=tools, employees=employees, today=datetime.now().strftime('%Y-%m-%d'))


# ==================== ВОЗВРАТ ИНСТРУМЕНТА ====================

@app.route('/return_tool', methods=['GET', 'POST'])
@login_required
def return_tool():
    conn = get_db()
    deadline_days = int(get_setting('deadline_days', '14'))
    deadline_date = (datetime.now() - timedelta(days=deadline_days)).strftime('%Y-%m-%d')

    if request.method == 'POST':
        transaction_id = request.form['transaction_id']
        return_date = request.form['return_date'] or datetime.now().strftime('%Y-%m-%d')

        transaction = conn.execute("SELECT * FROM transactions WHERE id = ? AND status = 'issued'", (transaction_id,)).fetchone()
        if not transaction:
            conn.close()
            flash('Ошибка: неверная операция возврата', 'danger')
            return redirect(url_for('return_tool'))

        # Обновляем транзакцию
        conn.execute('''
            UPDATE transactions
            SET return_date = ?, returned_to = ?, status = 'returned'
            WHERE id = ?
        ''', (return_date, session['user_id'], transaction_id))

        # Возвращаем статус инструмента
        conn.execute("UPDATE tools SET status = 'ok' WHERE id = ?",
                    (transaction['tool_id'],))

        conn.commit()
        conn.close()
        flash('Инструмент успешно возвращен на склад!', 'success')
        return redirect(url_for('dashboard'))

    # GET - показываем список выданных
    active = conn.execute('''
        SELECT tr.id, t.name as tool_name, t.inventory_number,
               e.fullname as employee_name, tr.issue_date,
               u.username as issued_by_name,
               CASE WHEN tr.issue_date < ? THEN 1 ELSE 0 END as is_overdue,
               CAST(julianday('now') - julianday(tr.issue_date) AS INTEGER) as days_issued
        FROM transactions tr
        JOIN tools t ON tr.tool_id = t.id
        JOIN employees e ON tr.employee_id = e.id
        JOIN users u ON tr.issued_by = u.id
        WHERE tr.status = 'issued'
        ORDER BY tr.issue_date DESC
    ''', (deadline_date,)).fetchall()

    conn.close()
    return render_template('return.html', active=active, today=datetime.now().strftime('%Y-%m-%d'),
                         deadline_days=deadline_days)


# ==================== РЕМОНТ ====================

@app.route('/repairs')
@login_required
def repairs_list():
    conn = get_db()
    active_repairs = conn.execute('''
        SELECT r.id, t.name as tool_name, t.inventory_number,
               r.sent_date, r.issue_description, r.repair_cost, r.notes
        FROM repairs r
        JOIN tools t ON r.tool_id = t.id
        WHERE r.return_date IS NULL
        ORDER BY r.sent_date DESC
    ''').fetchall()

    completed_repairs = conn.execute('''
        SELECT r.id, t.name as tool_name, t.inventory_number,
               r.sent_date, r.return_date, r.issue_description, r.repair_cost
        FROM repairs r
        JOIN tools t ON r.tool_id = t.id
        WHERE r.return_date IS NOT NULL
        ORDER BY r.return_date DESC
        LIMIT 20
    ''').fetchall()

    tools = conn.execute("SELECT * FROM tools WHERE status = 'ok' ORDER BY name").fetchall()
    conn.close()
    return render_template('repairs.html',
                         active_repairs=active_repairs,
                         completed_repairs=completed_repairs,
                         tools=tools,
                         today=datetime.now().strftime('%Y-%m-%d'))


@app.route('/send_to_repair', methods=['POST'])
@login_required
def send_to_repair():
    tool_id = request.form['tool_id']
    issue_description = request.form['issue_description']
    sent_date = request.form['sent_date'] or datetime.now().strftime('%Y-%m-%d')
    conn = get_db()
    conn.execute('INSERT INTO repairs (tool_id, sent_date, issue_description) VALUES (?, ?, ?)',
                (tool_id, sent_date, issue_description))
    conn.execute("UPDATE tools SET status = 'repair' WHERE id = ?", (tool_id,))
    conn.commit()
    conn.close()
    flash('Инструмент отправлен в ремонт', 'success')
    return redirect(url_for('repairs_list'))


@app.route('/return_from_repair/<int:repair_id>', methods=['POST'])
@login_required
def return_from_repair(repair_id):
    return_date = request.form['return_date'] or datetime.now().strftime('%Y-%m-%d')
    repair_cost = request.form.get('repair_cost', 0)
    conn = get_db()
    repair = conn.execute("SELECT * FROM repairs WHERE id = ?", (repair_id,)).fetchone()
    if repair:
        conn.execute("UPDATE repairs SET return_date = ?, repair_cost = ? WHERE id = ?",
                    (return_date, repair_cost, repair_id))
        conn.execute("UPDATE tools SET status = 'ok' WHERE id = ?", (repair['tool_id'],))
    conn.commit()
    conn.close()
    flash('Инструмент возвращен из ремонта', 'success')
    return redirect(url_for('repairs_list'))


# ==================== ЖУРНАЛ ====================

@app.route('/journal')
@login_required
def journal():
    conn = get_db()
    deadline_days = int(get_setting('deadline_days', '14'))
    deadline_date = (datetime.now() - timedelta(days=deadline_days)).strftime('%Y-%m-%d')

    status_filter = request.args.get('status', 'all')
    date_from = request.args.get('date_from', '')
    date_to = request.args.get('date_to', '')
    employee_filter = request.args.get('employee_id', '')

    query = '''
        SELECT tr.id, t.name as tool_name, t.inventory_number,
               e.fullname as employee_name, tr.issue_date, tr.return_date,
               u1.username as issued_by_name,
               u2.username as returned_to_name,
               tr.status,
               CASE WHEN tr.status = 'issued' AND tr.issue_date < ? THEN 1 ELSE 0 END as is_overdue,
               CAST(julianday('now') - julianday(tr.issue_date) AS INTEGER) as days_issued
        FROM transactions tr
        JOIN tools t ON tr.tool_id = t.id
        JOIN employees e ON tr.employee_id = e.id
        JOIN users u1 ON tr.issued_by = u1.id
        LEFT JOIN users u2 ON tr.returned_to = u2.id
        WHERE 1=1
    '''

    params = [deadline_date]

    if status_filter == 'issued':
        query += " AND tr.status = 'issued'"
    elif status_filter == 'returned':
        query += " AND tr.status = 'returned'"
    if date_from:
        query += " AND tr.issue_date >= ?"
        params.append(date_from)
    if date_to:
        query += " AND tr.issue_date <= ?"
        params.append(date_to)
    if employee_filter:
        query += " AND tr.employee_id = ?"
        params.append(employee_filter)

    query += " ORDER BY tr.issue_date DESC LIMIT 100"

    transactions = conn.execute(query, params).fetchall()
    employees = conn.execute("SELECT * FROM employees WHERE is_active = 1 ORDER BY fullname").fetchall()
    conn.close()
    return render_template('journal.html',
                         transactions=transactions,
                         employees=employees,
                         status_filter=status_filter,
                         date_from=date_from,
                         date_to=date_to,
                         employee_filter=employee_filter,
                         deadline_days=deadline_days)


# ==================== ОТЧЁТЫ ====================

@app.route('/reports')
@login_required
def reports():
    return render_template('reports.html')


@app.route('/report/tools_in_use')
@login_required
def report_tools_in_use():
    conn = get_db()
    tools = conn.execute('''
        SELECT t.name as instrument, t.inventory_number as inv_number,
               e.fullname as employee, e.department,
               tr.issue_date,
               CAST(julianday('now') - julianday(tr.issue_date) AS INTEGER) as days_issued
        FROM transactions tr
        JOIN tools t ON tr.tool_id = t.id
        JOIN employees e ON tr.employee_id = e.id
        WHERE tr.status = 'issued'
        ORDER BY tr.issue_date
    ''').fetchall()
    conn.close()

    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    wb = Workbook()
    ws = wb.active
    ws.title = "Инструменты на руках"

    # Заголовки
    headers = ['Инструмент', 'Инвентарный номер', 'Сотрудник', 'Отдел', 'Дата выдачи', 'Дней на руках']
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(name='Arial', bold=True, color="FFFFFF", size=11)
    thin_border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = thin_border

    # Данные
    for row_idx, row in enumerate(tools, 2):
        values = [row['instrument'], row['inv_number'], row['employee'],
                  row['department'], row['issue_date'], row['days_issued']]
        for col_idx, value in enumerate(values, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.font = Font(name='Arial', size=10)
            cell.border = thin_border
            if col_idx == 6:  # Дней на руках
                cell.alignment = Alignment(horizontal='center')

    # Автоширина
    for col in ws.columns:
        max_length = 0
        col_letter = col[0].column_letter
        for cell in col:
            if cell.value:
                max_length = max(max_length, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = min(max_length + 4, 50)

    # Сохраняем
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    return send_file(output,
                    mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                    as_attachment=True,
                    download_name=f'instruments_in_use_{datetime.now().strftime("%Y%m%d")}.xlsx')


@app.route('/report/employee_history')
@login_required
def report_employee_history():
    conn = get_db()
    history = conn.execute('''
        SELECT e.fullname as employee, e.department,
               t.name as instrument, t.inventory_number as inv_number,
               tr.issue_date, tr.return_date,
               CASE WHEN tr.status = 'issued' THEN 'На руках' ELSE 'Возвращён' END as status_text
        FROM transactions tr
        JOIN tools t ON tr.tool_id = t.id
        JOIN employees e ON tr.employee_id = e.id
        ORDER BY e.fullname, tr.issue_date DESC
    ''').fetchall()
    conn.close()

    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    wb = Workbook()
    ws = wb.active
    ws.title = "История по сотрудникам"

    headers = ['Сотрудник', 'Отдел', 'Инструмент', 'Инвентарный номер', 'Дата выдачи', 'Дата возврата', 'Статус']
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(name='Arial', bold=True, color="FFFFFF", size=11)
    thin_border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = thin_border

    for row_idx, row in enumerate(history, 2):
        values = [row['employee'], row['department'], row['instrument'],
                  row['inv_number'], row['issue_date'],
                  row['return_date'] or 'Не возвращено', row['status_text']]
        for col_idx, value in enumerate(values, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.font = Font(name='Arial', size=10)
            cell.border = thin_border

    for col in ws.columns:
        max_length = 0
        col_letter = col[0].column_letter
        for cell in col:
            if cell.value:
                max_length = max(max_length, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = min(max_length + 4, 50)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    return send_file(output,
                    mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                    as_attachment=True,
                    download_name=f'employee_history_{datetime.now().strftime("%Y%m%d")}.xlsx')


@app.route('/report/tool_history/<int:tool_id>')
@login_required
def report_tool_history(tool_id):
    conn = get_db()
    events = []

    transactions = conn.execute('''
        SELECT 'issue' as type, tr.issue_date as date,
               e.fullname as details, tr.status
        FROM transactions tr
        JOIN employees e ON tr.employee_id = e.id
        WHERE tr.tool_id = ?
    ''', (tool_id,)).fetchall()

    for t in transactions:
        events.append({
            'type': 'issue',
            'date': t['date'],
            'details': f"Выдан: {t['details']}",
            'status': t['status']
        })

    returns = conn.execute('''
        SELECT 'return' as type, tr.return_date as date,
               e.fullname as details, tr.status
        FROM transactions tr
        JOIN employees e ON tr.employee_id = e.id
        WHERE tr.tool_id = ? AND tr.return_date IS NOT NULL
    ''', (tool_id,)).fetchall()

    for r in returns:
        events.append({
            'type': 'return',
            'date': r['date'],
            'details': f"Возвращён: {r['details']}",
            'status': r['status']
        })

    repairs_sent = conn.execute('''
        SELECT 'repair_send' as type, r.sent_date as date,
               r.issue_description as details
        FROM repairs r
        WHERE r.tool_id = ?
    ''', (tool_id,)).fetchall()

    for rs in repairs_sent:
        events.append({
            'type': 'repair_send',
            'date': rs['date'],
            'details': f"Неисправность: {rs['details']}",
            'status': 'repair'
        })

    repairs_return = conn.execute('''
        SELECT 'repair_return' as type, r.return_date as date,
               r.repair_cost as details
        FROM repairs r
        WHERE r.tool_id = ? AND r.return_date IS NOT NULL
    ''', (tool_id,)).fetchall()

    for rr in repairs_return:
        events.append({
            'type': 'repair_return',
            'date': rr['date'],
            'details': f"Стоимость ремонта: {rr['details']} ₽",
            'status': 'returned'
        })

    events.sort(key=lambda x: x['date'], reverse=True)

    tools = conn.execute("SELECT * FROM tools ORDER BY name").fetchall()
    conn.close()
    return render_template('tool_history.html',
                         history=events,
                         tools=tools,
                         selected_tool=tool_id)

# ==================== НАСТРОЙКИ ====================

@app.route('/settings', methods=['GET', 'POST'])
@admin_required
def settings():
    if request.method == 'POST':
        deadline_days = request.form['deadline_days']
        conn = get_db()
        conn.execute("UPDATE settings SET value = ? WHERE key = 'deadline_days'", (deadline_days,))
        conn.commit()
        conn.close()
        flash('Настройки сохранены', 'success')
        return redirect(url_for('settings'))

    current_deadline = get_setting('deadline_days', '14')
    return render_template('settings.html', deadline_days=current_deadline)


# ==================== СПРАВОЧНИКИ ====================

@app.route('/tools_directory')
@login_required
def tools_directory():
    conn = get_db()
    tools = conn.execute('''
        SELECT t.*, 
               CASE WHEN t.status = 'issued' THEN e.fullname ELSE NULL END as issued_to
        FROM tools t
        LEFT JOIN transactions tr ON t.id = tr.tool_id AND tr.status = 'issued'
        LEFT JOIN employees e ON tr.employee_id = e.id
        ORDER BY t.inventory_number
    ''').fetchall()
    conn.close()
    return render_template('tools_directory.html', tools=tools)


@app.route('/add_tool', methods=['POST'])
@admin_required
def add_tool():
    name = request.form['name']
    inventory_number = request.form['inventory_number']
    description = request.form.get('description', '')

    conn = get_db()
    try:
        conn.execute('INSERT INTO tools (name, inventory_number, description, status) VALUES (?, ?, ?, ?)',
                    (name, inventory_number, description, 'ok'))
        conn.commit()
        flash('Инструмент добавлен', 'success')
    except sqlite3.IntegrityError:
        flash('Ошибка: такой инвентарный номер уже существует!', 'danger')
    finally:
        conn.close()

    return redirect(url_for('tools_directory'))


@app.route('/edit_tool/<int:tool_id>', methods=['POST'])
@admin_required
def edit_tool(tool_id):
    name = request.form['name']
    description = request.form.get('description', '')
    status = request.form['status']

    conn = get_db()
    conn.execute("UPDATE tools SET name = ?, description = ?, status = ? WHERE id = ?",
                (name, description, status, tool_id))
    conn.commit()
    conn.close()
    flash('Инструмент обновлен', 'success')
    return redirect(url_for('tools_directory'))


@app.route('/delete_tool/<int:tool_id>')
@admin_required
def delete_tool(tool_id):
    conn = get_db()
    issued = conn.execute("SELECT COUNT(*) FROM transactions WHERE tool_id = ? AND status = 'issued'", (tool_id,)).fetchone()[0]
    in_repair = conn.execute("SELECT COUNT(*) FROM repairs WHERE tool_id = ? AND return_date IS NULL", (tool_id,)).fetchone()[0]

    if issued > 0:
        conn.close()
        flash('Нельзя удалить инструмент: он сейчас на руках', 'danger')
        return redirect(url_for('tools_directory'))

    if in_repair > 0:
        conn.close()
        flash('Нельзя удалить инструмент: он находится в ремонте', 'danger')
        return redirect(url_for('tools_directory'))

    conn.execute("DELETE FROM tools WHERE id = ?", (tool_id,))
    conn.commit()
    conn.close()
    flash('Инструмент удалён', 'success')
    return redirect(url_for('tools_directory'))


@app.route('/employees_directory')
@login_required
def employees_directory():
    conn = get_db()
    employees = conn.execute("SELECT * FROM employees ORDER BY fullname").fetchall()
    conn.close()
    return render_template('employees_directory.html', employees=employees)


@app.route('/add_employee', methods=['POST'])
@admin_required
def add_employee():
    fullname = request.form['fullname']
    department = request.form.get('department', '')
    conn = get_db()
    conn.execute("INSERT INTO employees (fullname, department) VALUES (?, ?)",
                (fullname, department))
    conn.commit()
    conn.close()
    flash('Сотрудник добавлен', 'success')
    return redirect(url_for('employees_directory'))


@app.route('/edit_employee/<int:employee_id>', methods=['POST'])
@admin_required
def edit_employee(employee_id):
    fullname = request.form['fullname']
    department = request.form.get('department', '')
    is_active = 1 if request.form.get('is_active') else 0
    conn = get_db()
    conn.execute("UPDATE employees SET fullname = ?, department = ?, is_active = ? WHERE id = ?",
                (fullname, department, is_active, employee_id))
    conn.commit()
    conn.close()
    flash('Данные сотрудника обновлены', 'success')
    return redirect(url_for('employees_directory'))


@app.route('/delete_employee/<int:employee_id>')
@admin_required
def delete_employee(employee_id):
    conn = get_db()
    issued = conn.execute("SELECT COUNT(*) FROM transactions WHERE employee_id = ? AND status = 'issued'",
                         (employee_id,)).fetchone()[0]
    if issued > 0:
        conn.close()
        flash('Нельзя удалить сотрудника: у него есть невозвращённые инструменты', 'danger')
        return redirect(url_for('employees_directory'))

    conn.execute("DELETE FROM employees WHERE id = ?", (employee_id,))
    conn.commit()
    conn.close()
    flash('Сотрудник удалён', 'success')
    return redirect(url_for('employees_directory'))


@app.route('/users_directory')
@admin_required
def users_directory():
    conn = get_db()
    users = conn.execute("SELECT id, username, role FROM users ORDER BY username").fetchall()
    conn.close()
    return render_template('users_directory.html', users=users)


@app.route('/add_user', methods=['POST'])
@admin_required
def add_user():
    username = request.form['username']
    password = request.form['password']
    role = request.form['role']
    conn = get_db()
    try:
        conn.execute("INSERT INTO users (username, password_hash, role) VALUES (?, ?, ?)",
                    (username, generate_password_hash(password), role))
        conn.commit()
        flash('Пользователь добавлен', 'success')
    except sqlite3.IntegrityError:
        flash('Ошибка: такой пользователь уже существует', 'danger')
    finally:
        conn.close()
    return redirect(url_for('users_directory'))


@app.route('/delete_user/<int:user_id>')
@admin_required
def delete_user(user_id):
    if user_id == session['user_id']:
        flash('Нельзя удалить самого себя', 'danger')
        return redirect(url_for('users_directory'))
    conn = get_db()
    conn.execute("DELETE FROM users WHERE id = ?", (user_id,))
    conn.commit()
    conn.close()
    flash('Пользователь удален', 'success')
    return redirect(url_for('users_directory'))


# ==================== ЗАПУСК ====================

if __name__ == '__main__':
    init_db()
    print("=" * 50)
    print("СИСТЕМА УЧЕТА ИНСТРУМЕНТА ЗАПУЩЕНА")
    print("Откройте в браузере: http://127.0.0.1:5000")
    print("Логин администратора: admin")
    print("Пароль администратора: admin")
    print("Логин кладовщика: storekeeper")
    print("Пароль кладовщика: 123")
    print("=" * 50)
    if __name__ == '__main__':
    app.run(debug=False, host='0.0.0.0', port=5000)